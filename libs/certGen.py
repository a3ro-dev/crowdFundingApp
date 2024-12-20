from docx import Document
import os
import stat
import platform

def generate_docx_with_shapes(template_path, output_dir, details):
    """
    Generates a .docx certificate by replacing placeholders with actual details.

    Args:
        template_path (str): Path to the .docx template.
        output_dir (str): Directory to save the generated certificate.
        details (dict): Dictionary containing placeholders and their replacements.

    Returns:
        str: Path to the generated .docx file.
    """
    doc = Document(template_path)

    # Replace text in paragraphs
    for paragraph in doc.paragraphs:
        for key, value in details.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)

    # Replace text in text boxes and shapes
    for shape in doc.element.xpath('//w:drawing//w:t'):
        for key, value in details.items():
            if key in shape.text:
                shape.text = shape.text.replace(key, value)

    # Create output directory if it doesn't exist
    os.makedirs(output_dir, exist_ok=True)

    # Generate the output filename
    user_name = details.get("{name}", "output").replace(" ", "_")
    output_docx_path = os.path.join(output_dir, f"{user_name}.docx")

    # Save the updated .docx file
    doc.save(output_docx_path)

    # Set the file to read-only
    if platform.system() == 'Windows':
        os.chmod(output_docx_path, stat.S_IREAD)
    else:
        os.chmod(output_docx_path, stat.S_IREAD | stat.S_IRGRP | stat.S_IROTH)

    return output_docx_path

# Coded with ❤️ by a3ro-dev
